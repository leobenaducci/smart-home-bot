"""A document with nothing in it must fail, not ship.

Reported from real use: the Profesor produced an empty .pptx, announced a
download link for it, and mentioned a PDF that was never created. The builders
were happy to write a file out of a payload they had not understood — a deck of
one title slide — and `create()` reported success, so the model had a link to
announce and no way to know anything was wrong.

Refusing is strictly better: the caller gets a message naming the missing key
and can retry. These tests pin that, plus the near-miss shapes that are now
accepted rather than being a silent blank.
"""

import importlib.util
import os
import tempfile

import pytest

SCRIPT = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(
    os.path.abspath(__file__)))), "nanobot", "skills", "document", "create_doc.py")


@pytest.fixture()
def doc(monkeypatch):
    monkeypatch.setenv("NANOBOT_WORKSPACE", tempfile.mkdtemp())
    spec = importlib.util.spec_from_file_location("create_doc_under_test", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    # Offline: quickchart/codecogs are on the internet and irrelevant here.
    mod._fetch_image = lambda url, dest: False
    return mod


REAL = [{"heading": "Fracciones", "text": "Una parte de un entero."}]


def _has(module):
    import importlib.util
    return importlib.util.find_spec(module) is not None


@pytest.mark.parametrize("fmt", ["docx", "pdf", "xlsx", "pptx", "html"])
@pytest.mark.parametrize("payload,because", [
    ({"title": "Las fracciones"}, "no sections at all"),
    ({"title": "Las fracciones", "sections": []}, "an empty section list"),
    ({"title": "X", "content": "todo el texto"}, "content under an unknown key"),
    ({"title": "X", "sections": [{"heading": "A"}, {"heading": "B"}]}, "headings only"),
])
def test_a_contentless_payload_is_refused(doc, fmt, payload, because):
    result = doc.create({**payload, "format": fmt, "filename": f"t.{fmt}"})
    assert "error" in result, f"{because} produced a file for {fmt}"
    assert "No file was generated" in result["error"]
    assert not os.path.exists(os.path.join(str(doc.DOCS_DIR), f"t.{fmt}")), \
        "refused, but a file was written anyway"


def test_the_error_names_the_key_to_send(doc):
    """The message is read by a model that has to fix its own call."""
    err = doc.create({"format": "pptx", "filename": "t.pptx", "title": "X"})["error"]
    assert "sections" in err and "heading" in err and "text" in err
    err2 = doc.create({"format": "pptx", "filename": "t.pptx",
                       "sections": [{"heading": "A"}]})["error"]
    for key in ("text", "items", "questions", "table"):
        assert key in err2


@pytest.mark.parametrize("payload,because", [
    ({"slides": [{"heading": "A", "text": "hola mundo"}]}, "`slides` for `sections`"),
    ({"sections": ["primera linea", "segunda linea"]}, "sections as plain strings"),
    ({"sections": [{"heading": "A", "bullets": ["uno", "dos"]}]}, "`bullets` for `items`"),
])
def test_near_miss_shapes_render_instead_of_failing(doc, payload, because):
    """These are unambiguous — nothing else in the schema is called `slides` or
    `bullets` — so accepting them beats a failed call and a retry. `sections`
    as strings used to raise halfway through and lose the whole document."""
    result = doc.create({**payload, "format": "pptx", "filename": "ok.pptx"})
    assert "error" not in result, f"{because}: {result.get('error')}"

    from pptx import Presentation
    text = " ".join(
        shape.text_frame.text
        for slide in Presentation(result["saved_to"]).slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    for word in ("hola mundo", "primera linea", "uno"):
        if word in str(payload):
            assert word in text, f"{because}: {word!r} never reached the file"


def test_a_real_payload_still_works_in_every_format(doc):
    for fmt in ("docx", "pdf", "xlsx", "pptx", "html"):
        if fmt == "pdf" and not doc._find_browser() and not _has("reportlab"):
            # A PDF is a printed web page with a reportlab fallback; a box with
            # neither cannot make one at all, and saying so is honest where
            # failing would just report the environment as a defect.
            pytest.skip("no browser and no reportlab: nothing here can render a PDF")
        r = doc.create({"format": fmt, "filename": f"g.{fmt}",
                        "title": "Guía", "sections": REAL})
        assert "error" not in r, f"{fmt}: {r.get('error')}"
        assert os.path.getsize(r["saved_to"]) > 0


def test_html_may_carry_its_markup_instead_of_sections(doc):
    """The designer's escape hatch must not be caught by the content guard."""
    r = doc.create({"format": "html", "filename": "p.html",
                    "html": "<h1>Hola</h1><p>Una página</p>"})
    assert "error" not in r, r.get("error")
    assert "Una página" in open(r["saved_to"], encoding="utf-8").read()


def test_the_download_link_only_exists_when_the_file_does(doc):
    """The reported failure was a link to a file that was never written."""
    ok = doc.create({"format": "pptx", "filename": "s.pptx",
                     "title": "T", "sections": REAL})
    assert ok["download_link"].startswith("[📥 Download s.pptx](download:media/")
    bad = doc.create({"format": "pptx", "filename": "s2.pptx", "title": "T"})
    assert "download_link" not in bad and "error" in bad


# --- html also comes back as a PDF -------------------------------------------
# A designed piece is written as HTML because that is the only format where the
# model controls type, colour and layout — and it is usually also wanted on
# paper. Rendering it with a real browser keeps the two identical.

PIECE = ('<!doctype html><html><head><meta charset="utf-8"><style>'
         '@page{size:letter;margin:0}'
         '.p{width:8.5in;min-height:11in;background:#c0392b;display:grid}'
         '*{print-color-adjust:exact}</style></head>'
         '<body><div class="p"><h1>Restar fracciones</h1></div></body></html>')


def test_a_designed_page_comes_back_with_its_pdf(doc):
    if not doc._find_browser():
        pytest.skip("no browser on this box; the fallback path is covered below")
    r = doc.create({"format": "html", "filename": "afiche.html", "html": PIECE})
    assert "error" not in r, r
    assert r["pdf_file"].endswith("afiche.pdf")
    pdf = os.path.join(str(doc.WORKSPACE), r["pdf_file"])
    assert open(pdf, "rb").read(5) == b"%PDF-"
    # Both links reach the caller, and the message carries them in order.
    assert r["download_link"] in r["message"]
    assert r["pdf_download_link"] in r["message"]
    assert r["message"].index(r["download_link"]) < r["message"].index(r["pdf_download_link"])


def test_render_pdf_false_leaves_it_html_only(doc):
    r = doc.create({"format": "html", "filename": "solo.html",
                    "render_pdf": False, "html": PIECE})
    assert "error" not in r
    assert "pdf_download_link" not in r and "pdf_error" not in r


def test_a_box_with_no_browser_still_delivers_the_page(doc, monkeypatch):
    """The HTML is the deliverable and already exists by then. A missing
    renderer must cost the PDF, not the piece."""
    monkeypatch.setattr(doc, "_find_browser", lambda: None)
    monkeypatch.setitem(__import__("sys").modules, "weasyprint", None)
    r = doc.create({"format": "html", "filename": "sinpdf.html", "html": PIECE})
    assert "error" not in r, r
    assert os.path.getsize(r["saved_to"]) > 0
    assert r.get("pdf_error"), "a failed render must be reported, not silent"
    assert "No se pudo generar el PDF" in r["message"]


@pytest.mark.parametrize("fmt", ["docx", "xlsx", "pptx"])
def test_the_other_formats_are_not_double_rendered(doc, fmt):
    r = doc.create({"format": fmt, "filename": f"d.{fmt}", "title": "T",
                    "sections": [{"heading": "A", "text": "x"}]})
    assert "error" not in r
    assert "pdf_download_link" not in r


# --- a PDF is a printed web page ---------------------------------------------
# reportlab understood heading/paragraph/bullet/table/image stacked in order and
# silently dropped everything else, so "hazme un PDF" came back looking like a
# memo however it was asked for. The HTML builder already handles every section
# type and accepts raw markup; chromium prints it as designed.

def test_a_pdf_is_rendered_from_html(doc):
    if not doc._find_browser():
        pytest.skip("no browser on this box; the fallback is covered below")
    r = doc.create({"format": "pdf", "filename": "guia.pdf", "title": "Guía",
                    "sections": REAL})
    assert "error" not in r, r
    assert r["pdf_renderer"] == "chromium", r.get("pdf_renderer")
    assert open(r["saved_to"], "rb").read(5) == b"%PDF-"


def test_the_page_it_came_from_is_kept_but_not_offered(doc):
    """One document, one link. Two links per file is how a reply ends up
    offering the guide and the answer key and pointing both at one file."""
    if not doc._find_browser():
        pytest.skip("no browser on this box")
    r = doc.create({"format": "pdf", "filename": "g2.pdf", "title": "Guía",
                    "sections": REAL})
    assert os.path.isfile(os.path.join(str(doc.WORKSPACE), r["html_file"]))
    assert r["message"].count("download:") == 1, r["message"]
    assert "pdf_download_link" not in r


def test_a_designed_piece_can_be_asked_for_as_pdf(doc):
    """The `html` escape hatch is content, and a PDF is now that page printed."""
    if not doc._find_browser():
        pytest.skip("no browser on this box")
    r = doc.create({"format": "pdf", "filename": "afiche.pdf", "html": PIECE})
    assert "error" not in r, r
    assert open(r["saved_to"], "rb").read(5) == b"%PDF-"


def test_a_page_break_reaches_the_page(doc):
    r = doc.create({"format": "html", "filename": "pb.html", "title": "T",
                    "sections": [{"heading": "A", "text": "uno"},
                                 {"heading": "B", "text": "dos", "page_break": True}]})
    html = open(r["saved_to"], encoding="utf-8").read()
    assert 'class="pb"' in html and "break-before:page" in html


def test_print_is_not_darkened_by_a_dark_environment(doc):
    """Headless chromium renders with whatever colour scheme the environment
    reports, and the dark block sits after the print one — unscoped, it won."""
    r = doc.create({"format": "html", "filename": "d.html", "title": "T",
                    "sections": REAL})
    html = open(r["saved_to"], encoding="utf-8").read()
    assert "@media screen and (prefers-color-scheme:dark)" in html


def test_a_box_with_no_browser_still_produces_a_pdf(doc, monkeypatch):
    """reportlab is the fallback, not the front door — and says so."""
    pytest.importorskip("reportlab")
    monkeypatch.setattr(doc, "_find_browser", lambda: None)
    monkeypatch.setitem(__import__("sys").modules, "weasyprint", None)
    r = doc.create({"format": "pdf", "filename": "plain.pdf", "title": "T",
                    "sections": REAL})
    assert "error" not in r, r
    assert r["pdf_renderer"] == "reportlab"
    assert r.get("pdf_degraded"), "a degraded render must be visible"


def test_no_browser_and_no_reportlab_reports_both(doc, monkeypatch):
    monkeypatch.setattr(doc, "_find_browser", lambda: None)
    monkeypatch.setitem(__import__("sys").modules, "weasyprint", None)
    monkeypatch.setitem(__import__("sys").modules, "reportlab", None)
    r = doc.create({"format": "pdf", "filename": "none.pdf", "title": "T",
                    "sections": REAL})
    assert "error" in r and "reportlab" in r["error"]


# --- html -> docx -------------------------------------------------------------
# "Dámelo en Word" after a designed piece. The design cannot survive — docx has
# no grid, no absolute positioning, no page-sized artwork — so what crosses is
# the content, editable, and the caveat is said out loud rather than left for
# whoever opens the file to discover.

DOC_HTML = (
    '<!doctype html><html><head><style>body{color:red}</style><title>T</title></head>'
    '<body><h1>Guía de fracciones</h1>'
    '<p>Una fracción es <strong>una parte</strong> de un <em>entero</em>.</p>'
    '<h2>Ejercicios</h2><ul><li>Simplifica 12/16</li><li>Simplifica 18/24</li></ul>'
    '<ol><li>Primero</li><li>Segundo</li></ol>'
    '<table><thead><tr><th>Figura</th><th>Lados</th></tr></thead>'
    '<tbody><tr><td>Triángulo</td><td>3</td></tr></tbody></table>'
    '<p>Fin.</p></body></html>')


@pytest.fixture()
def converted(doc):
    from docx import Document
    r = doc.create({"format": "docx", "filename": "conv.docx", "html": DOC_HTML})
    assert "error" not in r, r
    return r, Document(r["saved_to"])


def test_a_page_can_be_asked_for_as_something_editable(converted):
    r, _ = converted
    assert r["from_html"] is True
    assert "download_link" in r


def test_the_structure_survives_even_though_the_design_does_not(converted):
    _, d = converted
    styled = [(p.style.name, p.text) for p in d.paragraphs if p.text.strip()]
    assert ("Title", "Guía de fracciones") in styled, styled
    assert ("Heading 1", "Ejercicios") in styled, styled
    assert [t for s, t in styled if s == "List Bullet"] == ["Simplifica 12/16", "Simplifica 18/24"]
    assert [t for s, t in styled if s == "List Number"] == ["Primero", "Segundo"]


def test_a_table_crosses_with_its_header(converted):
    _, d = converted
    assert len(d.tables) == 1
    assert [[c.text for c in row.cells] for row in d.tables[0].rows] == [
        ["Figura", "Lados"], ["Triángulo", "3"]]


def test_bold_and_italic_survive_as_runs(converted):
    _, d = converted
    runs = [(r.text, bool(r.bold), bool(r.italic))
            for p in d.paragraphs for r in p.runs if r.text.strip()]
    assert ("una parte", True, False) in runs, runs
    assert ("entero", False, True) in runs, runs


def test_css_and_scripts_do_not_become_text(converted):
    _, d = converted
    body = "\n".join(p.text for p in d.paragraphs)
    assert "color:red" not in body and "<" not in body


def test_an_embedded_image_crosses_and_leaves_no_temp_file(doc):
    import base64 as b64
    png = ("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8"
           "BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")
    r = doc.create({"format": "docx", "filename": "i.docx",
                    "html": f'<h1>T</h1><img src="data:image/png;base64,{png}"><p>x</p>'})
    assert "error" not in r, r
    from docx import Document
    assert len(Document(r["saved_to"]).inline_shapes) == 1
    assert not [f for f in os.listdir(str(doc.DOCS_DIR)) if f.startswith("_htmlimg")], \
        "the decoded image was left behind in media/"
    _ = b64


def test_the_caveat_is_said_out_loud(converted):
    r, _ = converted
    assert "design is not preserved" in r["note"]
    assert r["note"] in r["message"], "the caller must relay it, so it has to be in the message"


def test_a_broken_image_costs_the_picture_not_the_document(doc):
    r = doc.create({"format": "docx", "filename": "b.docx",
                    "html": '<h1>T</h1><img src="data:image/png;base64,!!!"><p>el texto</p>'})
    assert "error" not in r, r
    from docx import Document
    assert "el texto" in "\n".join(p.text for p in Document(r["saved_to"]).paragraphs)


# Everything below was found by review, and every one of them was invisible to
# the tests above because DOC_HTML happens not to contain the shape that breaks
# it. The first is the one that matters most: it made the feature's own primary
# case — convert the page you just built — return an empty document and report
# success.

def test_the_page_this_very_file_produces_converts(doc):
    """`_build_html` writes `<meta charset>` and `<meta name=viewport>`, and a
    void element never gets an end tag. With `meta` on the skip list the skip
    counter went up and never came down, so every word after the head vanished
    and `create()` cheerfully reported a document."""
    page = doc.create({"format": "html", "filename": "src.html", "title": "Guía",
                       "sections": [{"heading": "Fracciones", "text": "contenido real"}]})
    r = doc.create({"format": "docx", "filename": "conv2.docx",
                    "html": open(page["saved_to"], encoding="utf-8").read()})
    assert "error" not in r, r
    from docx import Document
    body = [p.text for p in Document(r["saved_to"]).paragraphs if p.text.strip()]
    assert body, "the document came out empty and said it had worked"
    assert "contenido real" in " ".join(body), body


@pytest.mark.parametrize("void", ["<meta charset='utf-8'>", "<link rel='x' href='y'>",
                                  "<br>", "<hr>", "<img src=''>", "<input>"])
def test_no_void_element_swallows_the_rest_of_the_page(doc, void):
    r = doc.create({"format": "docx", "filename": "v.docx",
                    "html": f"<p>antes</p>{void}<p>después</p>"})
    assert "error" not in r, r
    from docx import Document
    body = " ".join(p.text for p in Document(r["saved_to"]).paragraphs)
    assert "antes" in body and "después" in body, (void, body)


def test_a_cell_without_a_row_is_a_table_not_a_lost_document(doc):
    """`html.parser` inserts no implicit tags, and a model writes `<table><td>`."""
    r = doc.create({"format": "docx", "filename": "notr.docx",
                    "html": "<table><td>uno</td><td>dos</td></table>"})
    assert "error" not in r, r
    from docx import Document
    assert [[c.text for c in row.cells] for row in Document(r["saved_to"]).tables[0].rows] \
        == [["uno", "dos"]]


def test_a_space_between_inline_tags_is_a_word_gap(doc):
    """The runs were each correct, so inspecting runs could never see this."""
    r = doc.create({"format": "docx", "filename": "sp.docx",
                    "html": "<p><strong>Hola</strong> <em>mundo</em> y mas</p>"})
    from docx import Document
    assert [p.text for p in Document(r["saved_to"]).paragraphs if p.text.strip()] \
        == ["Hola mundo y mas"]


def test_a_nested_table_does_not_close_the_outer_one(doc):
    r = doc.create({"format": "docx", "filename": "nest.docx", "html": (
        "<table><tr><td>outer1<table><tr><td>inner</td></tr></table></td>"
        "<td>outer2</td></tr><tr><td>fila2a</td><td>fila2b</td></tr></table><p>fin</p>")})
    assert "error" not in r, r
    from docx import Document
    d = Document(r["saved_to"])
    outer = [[c.text for c in row.cells] for row in d.tables[0].rows]
    assert outer == [["outer1", "outer2"], ["fila2a", "fila2b"]], outer
    body = " ".join(p.text for p in d.paragraphs if p.text.strip())
    assert body == "fin", f"the outer table spilled into the body: {body!r}"


def test_an_image_in_a_cell_stays_in_the_cell(doc):
    png = ("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8"
           "BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")
    r = doc.create({"format": "docx", "filename": "cellimg.docx", "html": (
        f'<table><tr><td><img src="data:image/png;base64,{png}"></td>'
        f'<td>b</td></tr></table><p>fin</p>')})
    from docx import Document
    d = Document(r["saved_to"])
    assert len(d.inline_shapes) == 1
    # Where the drawing actually sits, not just what text is around it: a
    # picture appended to the body adds a paragraph with no text at all, so
    # reading paragraph text can never see it move.
    cell = d.tables[0].rows[0].cells[0]
    assert "graphicData" in cell._element.xml, "the picture is not in the cell"
    assert not any("graphicData" in p._element.xml for p in d.paragraphs), \
        "the picture was relocated into the body, out of order"


def test_a_line_break_inside_a_cell_survives(doc):
    r = doc.create({"format": "docx", "filename": "brcell.docx",
                    "html": "<table><tr><td>a<br>b</td></tr></table>"})
    from docx import Document
    assert Document(r["saved_to"]).tables[0].rows[0].cells[0].text == "a\nb"


def test_a_title_is_used_when_the_markup_has_none(doc):
    r = doc.create({"format": "docx", "filename": "ti.docx", "title": "Mi título",
                    "html": "<p>solo texto</p>"})
    from docx import Document
    body = [p.text for p in Document(r["saved_to"]).paragraphs if p.text.strip()]
    assert body == ["Mi título", "solo texto"], body


def test_the_markups_own_heading_wins_over_title(doc):
    r = doc.create({"format": "docx", "filename": "ti2.docx", "title": "Ignorado",
                    "html": "<h1>El de la página</h1><p>x</p>"})
    from docx import Document
    body = [p.text for p in Document(r["saved_to"]).paragraphs if p.text.strip()]
    assert body[0] == "El de la página" and "Ignorado" not in body, body


def test_sending_both_html_and_sections_says_which_one_was_used(doc):
    r = doc.create({"format": "docx", "filename": "both.docx", "html": "<p>x</p>",
                    "sections": [{"heading": "A", "text": "y"}]})
    assert "`sections` was ignored" in r["note"], r["note"]


def test_sections_still_take_the_normal_path(doc):
    """`html` is the exception; a section list is still the better docx."""
    r = doc.create({"format": "docx", "filename": "s.docx", "title": "Guía",
                    "sections": REAL})
    assert "from_html" not in r and "note" not in r
    from docx import Document
    assert "Fracciones" in "\n".join(p.text for p in Document(r["saved_to"]).paragraphs)


# --- two documents in one turn are two documents ------------------------------

def test_a_second_document_never_lands_on_the_first(doc):
    """Reported: "ambos links apuntan a la pauta". A prueba and its pauta are
    two calls, and nothing stopped the second overwriting the first — one file,
    written twice, offered as two."""
    a = doc.create({"format": "docx", "filename": "t.docx", "title": "Tarea",
                    "sections": [{"heading": "Tarea", "text": "el enunciado"}]})
    b = doc.create({"format": "docx", "filename": "t.docx", "title": "Pauta",
                    "sections": [{"heading": "Pauta", "text": "las respuestas"}]})
    assert a["download_link"] != b["download_link"], "two calls, one link"
    assert a["file"] != b["file"]
    assert os.path.isfile(a["saved_to"]) and os.path.isfile(b["saved_to"])
    from docx import Document
    assert "el enunciado" in "\n".join(p.text for p in Document(a["saved_to"]).paragraphs), \
        "the first document was overwritten by the second"


def test_the_renamed_file_keeps_its_extension(doc):
    a = doc.create({"format": "xlsx", "filename": "datos.xlsx", "title": "T",
                    "sections": REAL})
    b = doc.create({"format": "xlsx", "filename": "datos.xlsx", "title": "T",
                    "sections": REAL})
    assert b["file"].endswith(".xlsx") and b["file"] != a["file"], (a["file"], b["file"])
